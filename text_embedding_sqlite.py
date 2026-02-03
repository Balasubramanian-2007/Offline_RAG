# =========================
# IMPORTS
# =========================
from pypdf import PdfReader
import docx
import re
import os
import shutil
import sqlite3
import faiss
import numpy as np

from sentence_transformers import SentenceTransformer
from fastapi import FastAPI, UploadFile, File, HTTPException
from pydantic import BaseModel
from ollama import chat, ChatResponse

# =========================
# FASTAPI APP
# =========================
app = FastAPI()

UPLOAD_DIRECTORY = "./uploads"
os.makedirs(UPLOAD_DIRECTORY, exist_ok=True)

# =========================
# DATABASE (SQLite)
# =========================
conn = sqlite3.connect("rag.db", check_same_thread=False)
cur = conn.cursor()

cur.execute("""
CREATE TABLE IF NOT EXISTS chunks (
    chunk_id INTEGER PRIMARY KEY AUTOINCREMENT,
    source_name TEXT,
    heading TEXT,
    content TEXT,
    pg_start INTEGER,
    pg_end INTEGER,
    deleted INTEGER DEFAULT 0
)
""")
conn.commit()

# =========================
# EMBEDDING MODEL
# =========================
model = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")

# =========================
# QUERY SCHEMA
# =========================
class QueryRequest(BaseModel):
    query: str

# =========================
# HEADING DETECTION
# =========================
def isheading(line: str):
    score = 0
    line = line.strip()
    if not line:
        return False

    if len(line) <= 80: score += 1
    if line.endswith(":"): score += 1
    if line.isupper(): score += 2
    if len(line.split()) <= 10: score += 2
    if not line.endswith("."): score += 1
    if line.lower().startswith(("the ", "a ", "an ", "and ")): score -= 2

    return score if score >= 4 else False

# =========================
# PDF INGESTION
# =========================
def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    current_section = None

    for page_no, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if not text:
            continue

        for line in text.split("\n"):
            clean = line.strip()
            if not clean:
                continue

            if isheading(clean):
                if current_section:
                    cur.execute("""
                        INSERT INTO chunks (source_name, heading, content, pg_start, pg_end)
                        VALUES (?, ?, ?, ?, ?)
                    """, (
                        current_section["source"],
                        current_section["heading"],
                        current_section["content"],
                        current_section["start"],
                        current_section["end"]
                    ))
                    conn.commit()

                current_section = {
                    "source": file_path,
                    "heading": clean,
                    "content": "",
                    "start": page_no,
                    "end": page_no
                }
            else:
                if not current_section:
                    current_section = {
                        "source": file_path,
                        "heading": "GENERAL",
                        "content": "",
                        "start": page_no,
                        "end": page_no
                    }

                current_section["content"] += " " + clean
                current_section["end"] = page_no

    if current_section:
        cur.execute("""
            INSERT INTO chunks (source_name, heading, content, pg_start, pg_end)
            VALUES (?, ?, ?, ?, ?)
        """, (
            current_section["source"],
            current_section["heading"],
            current_section["content"],
            current_section["start"],
            current_section["end"]
        ))
        conn.commit()

# =========================
# WORD INGESTION
# =========================
def extract_text_from_word(file_path):
    doc = docx.Document(file_path)
    current_section = None

    for i, para in enumerate(doc.paragraphs):
        clean = para.text.strip()
        if not clean:
            continue

        if isheading(clean):
            if current_section:
                cur.execute("""
                    INSERT INTO chunks (source_name, heading, content, pg_start, pg_end)
                    VALUES (?, ?, ?, ?, ?)
                """, (
                    current_section["source"],
                    current_section["heading"],
                    current_section["content"],
                    current_section["start"],
                    current_section["end"]
                ))
                conn.commit()

            current_section = {
                "source": file_path,
                "heading": clean,
                "content": "",
                "start": i,
                "end": i
            }
        else:
            if not current_section:
                current_section = {
                    "source": file_path,
                    "heading": "GENERAL",
                    "content": "",
                    "start": i,
                    "end": i
                }

            current_section["content"] += " " + clean
            current_section["end"] = i

    if current_section:
        cur.execute("""
            INSERT INTO chunks (source_name, heading, content, pg_start, pg_end)
            VALUES (?, ?, ?, ?, ?)
        """, (
            current_section["source"],
            current_section["heading"],
            current_section["content"],
            current_section["start"],
            current_section["end"]
        ))
        conn.commit()

# =========================
# FILE UPLOAD ENDPOINT
# =========================
@app.post("/upload")
def upload_file(file: UploadFile = File(...)):
    file_path = os.path.join(UPLOAD_DIRECTORY, file.filename)

    with open(file_path, "wb") as f:
        shutil.copyfileobj(file.file, f)

    if file.filename.endswith(".pdf"):
        extract_text_from_pdf(file_path)
    elif file.filename.endswith(".docx"):
        extract_text_from_word(file_path)
    else:
        raise HTTPException(status_code=400, detail="Only PDF or DOCX allowed")

    # Fetch texts for embedding
    cur.execute("SELECT content FROM chunks WHERE deleted = 0")
    texts = [row[0] for row in cur.fetchall()]

    embeddings = model.encode(texts, normalize_embeddings=True).astype("float32")

    dim = embeddings.shape[1]
    index = faiss.IndexFlatL2(dim)

    if os.path.exists("faiss.index"):
        index = faiss.read_index("faiss.index")

    index.add(embeddings)
    faiss.write_index(index, "faiss.index")

    return {"message": "File indexed successfully"}

# =========================
# FETCH CHUNKS
# =========================
def fetch_chunks(chunk_ids):
    placeholders = ",".join(["?"] * len(chunk_ids))
    query = f"""
        SELECT heading, content
        FROM chunks
        WHERE chunk_id IN ({placeholders}) AND deleted = 0
    """
    cur.execute(query, chunk_ids)
    rows = cur.fetchall()

    context = ""
    for h, c in rows:
        context += f"\n[SECTION: {h}]\n{c}\n"

    return context.strip()

# =========================
# PROMPT
# =========================
def build_prompt(context, question):
    return f"""
You are a helpful technical assistant.
Answer ONLY from the context below.
If not found, say "Not found in the document".

Context:
{context}

Question:
{question}

Answer:
""".strip()

# =========================
# OLLAMA CALL
# =========================
def ask_ollama(prompt):
    response: ChatResponse = chat(
        model="qwen2.5:1.5b",
        messages=[{"role": "user", "content": prompt}]
    )
    return response.message.content

# =========================
# QUERY ENDPOINT
# =========================

@app.post("/debug")
def debug(req:dict):
    return req
@app.post("/result")
def query_rag(req: QueryRequest):
    embedding = model.encode(req.query, normalize_embeddings=True)
    embedding = np.array(embedding).astype("float32").reshape(1, -1)

    if not os.path.exists("faiss.index"):
        raise HTTPException(status_code=404, detail="FAISS index not found")

    index = faiss.read_index("faiss.index")
    _, I = index.search(embedding, 5)

    chunk_ids = [i + 1 for i in I[0].tolist()]
    context = fetch_chunks(chunk_ids)

    prompt = build_prompt(context, req.query)
    answer = ask_ollama(prompt)

    return {
        "question": req.query,
        "answer": answer
    }